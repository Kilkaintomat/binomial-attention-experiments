"""Run a small, explicitly labeled pointer reimplementation and workbook audit."""
import csv
import json
import time
from pathlib import Path
import numpy as np
import torch
from openpyxl import load_workbook
from transport_reasoning.config import ROOT,write_json
from transport_reasoning.data.pointer import Pointer,views,generate,metrics
from transport_reasoning.filters import coefficients,moments
from transport_reasoning.reporting import table

torch.set_num_threads(2)
rows=[]
for seed in (11,23,37):
    qview,kview=views(seed)
    for method in ("binomial","finite_geometric","exponential_truncated"):
        destination=ROOT/"bridge"/method/f"seed_{seed}"; destination.mkdir(parents=True,exist_ok=True)
        if (destination/"results.json").exists():
            rows.extend(json.loads((destination/"results.json").read_text())); continue
        model=Pointer(seed); optimizer=torch.optim.Adam(model.parameters(),lr=.02)
        gen=torch.Generator().manual_seed(seed+7003)
        start=time.perf_counter(); curve=[]
        for step in range(1,501):
            depth=2+(step-1)%2
            data=generate(32,depth,4,gen,qview,kview)
            y,_=model(data,method,6,4); labels=data[4]
            correct=y.gather(-1,labels[:,:,None]).squeeze(-1)
            loss=-(correct/y.sum(-1).clamp_min(1e-30)).clamp_min(1e-30).log().mean()
            optimizer.zero_grad(); loss.backward(); torch.nn.utils.clip_grad_norm_(model.parameters(),5); optimizer.step()
            if step in (5,10,20,25,30,40,50,75,100,200,300,500): curve.append({"step":step,"loss":float(loss)})
        elapsed=time.perf_counter()-start; run=[]
        with torch.no_grad():
            for condition,depth,width,hops in (("iid_depth2",2,4,6),("iid_depth3",3,4,6),("depth6_fixed_budget",6,4,6),("depth6_scaled_budget",6,4,12)):
                gen=torch.Generator().manual_seed(seed+depth*117+width*23)
                values=[]
                for _ in range(16): values.append(metrics(model,generate(64,depth,width,gen,qview,kview),method,hops,width))
                row={"implementation":"reimplementation","seed":seed,"method":method,"condition":condition,"H":hops,
                     "training_seconds":elapsed,"training_steps":500,"train_hops":6,"parameters":2048,
                     **{k:float(np.mean([v[k] for v in values])) for k in values[0]},**moments(coefficients(method,hops))}
                run.append(row)
        torch.save({"model":model.state_dict(),"optimizer":optimizer.state_dict(),"data_rng":gen.get_state()},destination/"checkpoint_last.pt")
        write_json(destination/"curve.json",curve); write_json(destination/"results.json",run); rows.extend(run)
table(ROOT/"summaries"/"pointer_bridge.csv",rows)
book=load_workbook('/inputs/binomial_attention_data.xlsx',read_only=True,data_only=True)
audit={"source":"attached workbook; historical results only","sheet_count":len(book.sheetnames),"sheets":{}}
for name in ("Final evaluations","Learning runs","Filter weights"):
    values=list(book[name].values)
    header_index=next(i for i,row in enumerate(values) if "seed" in [str(v).lower() for v in row]) if name!="Filter weights" else None
    if header_index is not None:
        header=[str(v).lower() if v is not None else None for v in values[header_index]]
        records=[dict(zip(header,row)) for row in values[header_index+1:] if any(v is not None for v in row)]
        audit["sheets"][name]={"columns":[str(v) for v in header if v is not None],"data_rows":len(records)}
        if name=="Final evaluations":
            numeric={}
            for row in records:
                method=row.get("method"); condition=row.get("evaluation condition")
                if method is None: continue
                key=str(method)+"/"+str(condition)
                for col,val in row.items():
                    if isinstance(val,(int,float)) and col!="seed": numeric.setdefault(key,{}) .setdefault(str(col),[]).append(val)
            audit["per_seed_recomputed_means"]={k:{c:float(np.mean(v)) for c,v in vals.items()} for k,vals in numeric.items()}
            expected={"binomial/depth6_fixed_budget":{"accuracy (fraction)":.9984537760416666,"terminal mass (fraction)":.0161506},
                      "finite_ppr/depth6_fixed_budget":{"accuracy (fraction)":.9999186197916666,"terminal mass (fraction)":.108668}}
            checks=[]
            for key,columns in expected.items():
                for column,value in columns.items():
                    measured=audit["per_seed_recomputed_means"][key][column]
                    checks.append({"key":key,"metric":column,"recomputed":measured,"plan_rounded_value":value,"matches":abs(measured-value)<1e-6})
            audit["plan_numeric_checks"]=checks
            if not all(item["matches"] for item in checks): raise ValueError("Workbook disagrees with plan")
book.close(); write_json(ROOT/"workbook_audit.json",audit)
write_json(ROOT/"pointer_bridge_complete.json",{"rows":len(rows),"status":"completed","exact_old_reproduction":False})
print(json.dumps({"pointer_rows":len(rows),"status":"completed"}))
